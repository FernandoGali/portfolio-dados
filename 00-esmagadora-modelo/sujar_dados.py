"""
Gera a versão SUJA da Esmagadora Modelo a partir de dados_limpos/.

Imita o que acontece em exportações reais de planilha e sistema:
CSV em padrão brasileiro (separador ;, decimal vírgula, encoding Windows-1252),
datas em formatos misturados, texto digitado de vários jeitos, zeros no lugar
de vazio, duplicados por reenvio, erros de unidade e de digitação.

Saída: dados_sujos/ (CSVs + apontamentos_mensais/*.xlsx)
       _gabarito/gabarito_sujeira.md (o que foi sujado e quantas linhas)
"""
from pathlib import Path
import unicodedata
import numpy as np
import pandas as pd
from openpyxl import Workbook

rng = np.random.default_rng(7)
BASE = Path(__file__).parent
LIMPO, SUJO, GAB = BASE / "dados_limpos", BASE / "dados_sujos", BASE / "_gabarito"
SUJO.mkdir(exist_ok=True)
(SUJO / "apontamentos_mensais").mkdir(exist_ok=True)
GAB.mkdir(exist_ok=True)

gabarito = []


def anota(tabela, problema, n, como_detectar):
    gabarito.append((tabela, problema, n, como_detectar))


def ler(nome):
    return pd.read_csv(LIMPO / f"{nome}.csv")


def amostra(df, frac):
    n = max(1, int(len(df) * frac))
    return rng.choice(df.index.to_numpy(), n, replace=False)


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def datas_misturadas(serie, tabela):
    d = pd.to_datetime(serie)
    out = d.dt.strftime("%d/%m/%Y").astype(object)
    sorteio = rng.random(len(d))
    m1 = sorteio < 0.10
    m2 = (sorteio >= 0.10) & (sorteio < 0.15)
    out[m1] = d[m1].dt.strftime("%Y-%m-%d")
    out[m2] = d[m2].dt.strftime("%d-%m-%y")
    anota(tabela, "Datas em 3 formatos (dd/mm/aaaa, aaaa-mm-dd, dd-mm-aa)",
          int(m1.sum() + m2.sum()), "Converter com formato explícito; contar o que vira nulo")
    return out


def variante_texto(s):
    """Uma forma 'digitada' diferente do mesmo texto."""
    op = rng.integers(0, 5)
    if op == 0:
        return s.upper()
    if op == 1:
        return s.lower()
    if op == 2:
        return sem_acento(s)
    if op == 3:
        return s + "  "
    return " " + s


def salvar(df, nome):
    df.to_csv(SUJO / f"{nome}.csv", index=False, sep=";", decimal=",", encoding="cp1252", errors="replace")


# ======================================================== dimensões (iguais, formato BR)
for nome in ["dim_calendario", "dim_motivo_parada", "dim_equipamento", "dim_fornecedor", "dim_cliente"]:
    salvar(ler(nome), nome)

# ======================================================== produção
p = ler("fato_producao")
falta = amostra(p, 0.012)
anota("fato_producao", "Dias inteiros faltando (boletim não lançado)", len(falta),
      "Comparar com dim_calendario (anti join)")
p = p.drop(index=falta)
dup = p.loc[amostra(p, 0.006)].copy()
dup["soja_esmagada_t"] = (dup["soja_esmagada_t"] * rng.uniform(0.97, 1.03, len(dup))).round(2)
p = pd.concat([p, dup]).sort_values("data", kind="stable")
anota("fato_producao", "Dias duplicados com valor retificado (2 versões do mesmo dia)", len(dup),
      "Contar linhas por data; decidir qual versão vale")
idx = amostra(p.reset_index(drop=True), 0.02)
p = p.reset_index(drop=True)
p.loc[idx, "energia_kwh"] = np.nan
anota("fato_producao", "energia_kwh vazia", len(idx), "isna / IS NULL")
idx = amostra(p[p.hexano_l > 0], 0.004)
p.loc[idx, "hexano_l"] = p.loc[idx, "hexano_l"] * 10
anota("fato_producao", "hexano_l digitado com um zero a mais (x10)", len(idx),
      "Índice L/t muito acima do normal (> 8 L/t)")
p["data"] = datas_misturadas(p["data"], "fato_producao")
salvar(p, "fato_producao")

# ======================================================== paradas
pa = ler("fato_paradas")
mot = ler("dim_motivo_parada").set_index("id_motivo")["motivo"]
eq = ler("dim_equipamento").set_index("id_equipamento")["tag"]
pa["motivo"] = pa["id_motivo"].map(mot)
pa["equipamento"] = pa["id_equipamento"].map(eq)
idx = amostra(pa, 0.25)
pa.loc[idx, "motivo"] = pa.loc[idx, "motivo"].map(variante_texto)
anota("fato_paradas", "Motivo digitado de formas diferentes (caixa, acento, espaços)", len(idx),
      "value_counts / GROUP BY motivo; padronizar e casar com dim_motivo_parada")
tem_eq = pa["equipamento"].notna()
idx = amostra(pa[tem_eq], 0.2)
pa.loc[idx, "equipamento"] = pa.loc[idx, "equipamento"].map(
    lambda t: rng.choice([t.replace("-0", ""), t.replace("-", " ").lower(), t.replace("-", "")]))
anota("fato_paradas", "Tag de equipamento sem padrão (DT-02, DT2, dt 02, DT02)", len(idx),
      "Remover hífen/espaço/zero à esquerda e comparar com dim_equipamento")
pa["turno"] = pa["turno"].astype(object)
idx = amostra(pa, 0.15)
pa.loc[idx, "turno"] = pa.loc[idx, "turno"].map(
    lambda t: rng.choice([{"I": "1", "II": "2", "III": "3"}[t], f"Turno {t}"]))
anota("fato_paradas", "Turno como I / 1 / Turno I", len(idx), "value_counts do turno")


def fmt_duracao(m):
    r = rng.random()
    if r < 0.80:
        return str(m)
    if r < 0.90:
        return f"{m // 60}h{m % 60:02d}"
    if r < 0.97:
        return f"{m // 60:02d}:{m % 60:02d}"
    return f"{m} min"


pa["duracao"] = pa["duracao_min"].map(fmt_duracao)
anota("fato_paradas", "Duração como texto em 4 formatos (80 / 1h20 / 01:20 / 80 min)",
      int((~pa["duracao"].str.fullmatch(r"\d+")).sum()), "Coluna vira texto; regex para converter em minutos")
dup = pa.loc[amostra(pa, 0.01)]
pa = pd.concat([pa, dup]).sort_values("id_parada", kind="stable")
anota("fato_paradas", "Eventos duplicados (mesmo id_parada)", len(dup), "Contar por id_parada")
pa = pa[["id_parada", "data", "hora_inicio", "turno", "motivo", "equipamento", "duracao", "descricao"]]
pa["data"] = datas_misturadas(pa["data"], "fato_paradas")
salvar(pa, "fato_paradas")

# ======================================================== laboratório do farelo
l = ler("fato_lab_farelo")
n0 = int(l["oleo_casca_pct"].isna().sum() + l["proteina_soluvel_pct"].isna().sum())
l["oleo_casca_pct"] = l["oleo_casca_pct"].fillna(0)
l["proteina_soluvel_pct"] = l["proteina_soluvel_pct"].fillna(0)
anota("fato_lab_farelo", "Análise não feita lançada como 0 (óleo na casca, proteína solúvel)", n0,
      "Zeros fisicamente impossíveis; trocar por nulo antes de calcular média")
idx = amostra(l, 0.003)
l.loc[idx, "proteina_pct"] = (l.loc[idx, "proteina_pct"] / 10).round(2)
anota("fato_lab_farelo", "Proteína com vírgula no lugar errado (4,59 em vez de 45,9)", len(idx),
      "Valores fora da faixa física (proteína < 40)")
idx = amostra(l, 0.002)
l.loc[idx, "umidade_pct"] = (l.loc[idx, "umidade_pct"] * 10).round(1)
anota("fato_lab_farelo", "Umidade sem vírgula (113 em vez de 11,3)", len(idx), "Umidade > 20")
l["turno"] = l["turno"].astype(object)
idx = amostra(l, 0.1)
l.loc[idx, "turno"] = l.loc[idx, "turno"].map({"I": "1", "II": "2", "III": "3"})
anota("fato_lab_farelo", "Turno como número em parte das linhas", len(idx), "value_counts do turno")
l["data"] = datas_misturadas(l["data"], "fato_lab_farelo")
salvar(l, "fato_lab_farelo")

# ======================================================== soja industrializada
ls = ler("fato_lab_soja_industrializada")
idx = amostra(ls, 0.03)
cols = ["proteina_pct", "oleo_pct", "umidade_pct", "impureza_pct", "avariados_pct", "esverdeados_pct"]
ls.loc[idx, cols] = 0.0
anota("fato_lab_soja_industrializada", "Dia sem análise com todas as colunas = 0,00 (igual ao boletim real)",
      len(idx), "Linha com todos os valores zero")
ls["data"] = datas_misturadas(ls["data"], "fato_lab_soja_industrializada")
salvar(ls, "fato_lab_soja_industrializada")

# ======================================================== recebimento
r = ler("fato_recebimento")
forn = ler("dim_fornecedor").set_index("id_fornecedor")
r["cod_fornecedor"] = r["id_fornecedor"].map(forn["cod_fornecedor"]).astype(object)
r["fornecedor"] = r["id_fornecedor"].map(forn["fornecedor"]).astype(object)
idx = amostra(r, 0.30)


def variante_forn(s):
    if rng.random() < 0.3:
        s = s.replace("Fazenda", "Faz.").replace("Agropecuária", "Agrop.").replace("Cooperativa Agrícola", "Coop. Agrícola")
    return variante_texto(s)


r.loc[idx, "fornecedor"] = r.loc[idx, "fornecedor"].map(variante_forn)
anota("fato_recebimento", "Nome do fornecedor com variações (abreviação, caixa, acento, espaços)", len(idx),
      "Agrupar por nome dá mais fornecedores do que existem; usar cod_fornecedor")
idx = amostra(r, 0.04)
r.loc[idx, "cod_fornecedor"] = np.nan
anota("fato_recebimento", "cod_fornecedor vazio (só o nome preenchido)", len(idx),
      "Recuperar o código pelo nome padronizado")
idx = amostra(r, 0.003)
r.loc[idx, "peso_liquido_t"] = r.loc[idx, "peso_liquido_t"] * 1000
anota("fato_recebimento", "Peso líquido lançado em kg em vez de t", len(idx),
      "Carga acima de 60 t é impossível para um caminhão")
idx = amostra(r, 0.003)
r.loc[idx, "umidade_pct"] = (r.loc[idx, "umidade_pct"] * 10).round(1)
anota("fato_recebimento", "Umidade sem vírgula (125 em vez de 12,5)", len(idx), "Umidade > 30")
idx = amostra(r, 0.015)
r.loc[idx, "impureza_pct"] = np.nan
anota("fato_recebimento", "Impureza não informada", len(idx), "isna; decidir se exclui ou imputa")
dup = r.loc[amostra(r, 0.006)]
r = pd.concat([r, dup]).sort_values("id_ticket", kind="stable")
anota("fato_recebimento", "Ticket duplicado por reenvio do sistema da balança", len(dup),
      "Contar por id_ticket")
r = r.drop(columns="id_fornecedor")
r = r[["id_ticket", "data", "hora_entrada", "cod_fornecedor", "fornecedor", "peso_bruto_t", "peso_liquido_t",
       "umidade_pct", "impureza_pct", "avariados_pct", "esverdeados_pct", "proteina_pct", "oleo_pct", "desconto_t"]]
r["data"] = datas_misturadas(r["data"], "fato_recebimento")
salvar(r, "fato_recebimento")

# ======================================================== expedição
e = ler("fato_expedicao")
cli = ler("dim_cliente").set_index("id_cliente")["cliente"]
e["cliente"] = e["id_cliente"].map(cli).astype(object)
idx = amostra(e, 0.2)
e.loc[idx, "cliente"] = e.loc[idx, "cliente"].map(variante_texto)
anota("fato_expedicao", "Nome do cliente com variações", len(idx), "Casar com dim_cliente pelo nome padronizado")
e["produto"] = e["produto"].astype(object)
idx = amostra(e, 0.12)
e.loc[idx, "produto"] = e.loc[idx, "produto"].map(
    lambda x: rng.choice(["FARELO", "Farelo de soja", "farelo"]) if x == "Farelo" else rng.choice(["OLEO", "Óleo degomado", "oleo"]))
anota("fato_expedicao", "Produto escrito de formas diferentes", len(idx), "value_counts do produto")
est = e.loc[amostra(e, 0.003)].copy()
est["quantidade_t"] = -est["quantidade_t"]
e = pd.concat([e, est]).sort_values("id_expedicao", kind="stable")
anota("fato_expedicao", "Estornos (mesma nota com quantidade negativa)", len(est),
      "quantidade_t < 0; somar por id_expedicao")
e = e.drop(columns="id_cliente")[["id_expedicao", "data", "produto", "cliente", "quantidade_t"]]
e["data"] = pd.to_datetime(e["data"]).dt.strftime("%d/%m/%Y").astype(object)
idx = amostra(e[e["data"].str.endswith("2026")], 0.002)
e.loc[idx, "data"] = e.loc[idx, "data"].str.replace("2026", "2062")
anota("fato_expedicao", "Ano digitado errado (2062 em vez de 2026)", len(idx), "Datas no futuro")
salvar(e, "fato_expedicao")

# ======================================================== apontamentos mensais (Power Query)
p_limpo = ler("fato_producao")
p_limpo["data"] = pd.to_datetime(p_limpo["data"])
cab_padrao = ["Data", "Soja esmagada (t)", "Horas trabalhadas", "Farelo (t)", "Óleo degomado (t)",
              "Hexano (L)", "Vapor (kg)", "Energia (kWh)"]
cab_alt = ["DATA", "Soja Esmagada t", "Hrs Trab.", "Farelo t", "Oleo t", "Hexano L", "Vapor kg", "Energia kWh"]
meses_alt = {4, 9}
for mes in range(1, 13):
    g = p_limpo[(p_limpo.data.dt.year == 2025) & (p_limpo.data.dt.month == mes)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletim"
    ws["A1"] = "BOLETIM DIÁRIO DE PRODUÇÃO - ESMAGADORA MODELO"
    ws.merge_cells("A1:H1")
    ws["A2"] = f"Mês de referência: {mes:02d}/2025"
    ws.append([])
    ws.append(cab_alt if mes in meses_alt else cab_padrao)
    for _, row in g.iterrows():
        ws.append([row.data.to_pydatetime(), row.soja_esmagada_t, row.horas_trabalhadas, row.farelo_produzido_t,
                   row.oleo_degomado_t, row.hexano_l, row.vapor_kg, row.energia_kwh])
    ws.append([])
    ws.append(["TOTAL", g.soja_esmagada_t.sum(), g.horas_trabalhadas.sum(), g.farelo_produzido_t.sum(),
               g.oleo_degomado_t.sum(), g.hexano_l.sum(), g.vapor_kg.sum(), g.energia_kwh.sum()])
    wb.save(SUJO / "apontamentos_mensais" / f"producao_2025_{mes:02d}.xlsx")
anota("apontamentos_mensais", "Título e linha vazia antes do cabeçalho (cabeçalho na linha 4)", 12,
      "Power Query: remover linhas superiores e promover cabeçalho")
anota("apontamentos_mensais", "Abril e setembro com nomes de coluna diferentes", 2,
      "O Acrescentar cria colunas duplicadas; renomear antes de combinar")
anota("apontamentos_mensais", "Linha de TOTAL no fim de cada arquivo", 12,
      "Filtrar a linha TOTAL senão a soma dobra")

# ======================================================== gabarito
linhas = ["# Gabarito da versão suja", "",
          "Não abra antes da Fase 7. Tente achar os problemas sozinho e depois confira aqui.", "",
          "| Tabela | Problema | Linhas afetadas | Como detectar |", "|---|---|---|---|"]
linhas += [f"| {t} | {p} | {n:,} | {c} |".replace(f"{n:,}", f"{n:,}".replace(",", ".")) for t, p, n, c in gabarito]
linhas += ["", "Formato de todos os CSVs sujos: separador `;`, decimal `,`, encoding Windows-1252 (ANSI)."]
(GAB / "gabarito_sujeira.md").write_text("\n".join(linhas), encoding="utf-8")
print("\n".join(linhas))
